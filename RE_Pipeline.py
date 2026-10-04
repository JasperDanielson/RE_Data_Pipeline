#!/usr/bin/env python3
"""Canadian multifamily research evidence pipeline, version 2.0.1.

Python 3.11+. Install: pip install pandas numpy requests matplotlib openpyxl
Run: python RE_Pipeline.py
Repeat without network: python RE_Pipeline.py --offline
Refresh sources: python RE_Pipeline.py --refresh

The default evidence window is 2019 through the selected rental report year.
Sources: CMHC archived Rental Market Survey workbooks; Statistics Canada
17-10-0148 (CMA population), 17-10-0009 (quarterly national population),
34-10-0154 (monthly CMA construction); Bank of Canada Valet JSON.

Design: descriptive evidence, explicit geography/period/quality flags, no
imputed source values, no arbitrary market ranking, no underpowered regression.
A calendar-year construction flow is kept distinct from October rental surveys
and July population estimates. Output publication occurs only after validation.
This standalone program uses public data sources and standard Python libraries.
"""
from __future__ import annotations

import argparse
import hashlib
import io
import json
import logging
import math
import os
from pathlib import Path
import platform
import re
import shutil
import sys
import time
from datetime import datetime, timezone
from zoneinfo import ZoneInfo
import zipfile

import numpy as np
import pandas as pd
import requests
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.ticker import MaxNLocator, FuncFormatter
from matplotlib.colors import TwoSlopeNorm
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

VERSION = "2.0.1"
CITIES = ["Toronto", "Ottawa", "Montreal", "Halifax", "Calgary", "Edmonton", "Vancouver"]
GEOS = ["Canada", "Canada CMAs", *CITIES]
DISPLAY = {"Canada": "Canada (10,000+ rental centres)", "Canada CMAs": "Canada CMA aggregate",
           "Ottawa": "Ottawa (Ontario part)", "Montreal": "Montréal"}
RMS_LABELS = {"Canada 10,000+": "Canada", "Canada CMAs": "Canada CMAs",
              "Toronto CMA": "Toronto", "Vancouver CMA": "Vancouver", "Calgary CMA": "Calgary",
              "Edmonton CMA": "Edmonton", "Montréal CMA": "Montreal", "Halifax CMA": "Halifax",
              "Ottawa-Gatineau CMA (Ont. part)": "Ottawa"}
POP_LABELS = {"Canada": "Canada", "All census metropolitan areas, Canada": "Canada CMAs",
              "Toronto (CMA), Ontario": "Toronto", "Vancouver (CMA), British Columbia": "Vancouver",
              "Calgary (CMA), Alberta": "Calgary", "Edmonton (CMA), Alberta": "Edmonton",
              "Montréal (CMA), Quebec": "Montreal", "Halifax (CMA), Nova Scotia": "Halifax",
              "Ottawa - Gatineau (CMA), Ontario part, Ontario": "Ottawa"}
CON_LABELS = {"Census metropolitan areas": "Canada CMAs", "Toronto, Ontario": "Toronto",
              "Vancouver, British Columbia": "Vancouver", "Calgary, Alberta": "Calgary",
              "Edmonton, Alberta": "Edmonton", "Montréal, Quebec": "Montreal",
              "Halifax, Nova Scotia": "Halifax", "Ottawa-Gatineau, Ontario part, Ontario/Quebec": "Ottawa"}
CON_MEASURES = {"Housing starts": "starts", "Housing completions": "completions",
                "Housing under construction": "units_under_construction"}
BOC_SERIES = {"overnight_rate": "V39079", "goc_5y": "BD.CDN.5YR.DQ.YLD", "goc_10y": "BD.CDN.10YR.DQ.YLD"}
RMS_BASE = ("https://eppd1strscr01.blob.core.windows.net/cmhcprdcontainer/sf/project/archive/"
            "data_tables/data_tables/rental_market_report_data_tables_canada/")
RMS_PAGE = "https://www.cmhc-schl.gc.ca/chic/Listing?item_ID=%7BA6C8DBDA-51BA-4EE7-9432-5D0522FE2A8D%7D"
BLUE, TEAL, ORANGE, GREY = "#173F5F", "#087F8C", "#B9572D", "#68717A"
LOG = logging.getLogger("RE_Pipeline")


def text(v):
    return "" if pd.isna(v) else re.sub(r"\s+", " ", str(v)).strip()


def json_safe(v):
    if isinstance(v, dict): return {str(k): json_safe(x) for k, x in v.items()}
    if isinstance(v, (list, tuple)): return [json_safe(x) for x in v]
    if isinstance(v, (pd.Timestamp, datetime)): return v.isoformat()
    if isinstance(v, Path): return str(v)
    if isinstance(v, np.generic): return json_safe(v.item())
    if v is None or v is pd.NA or v is pd.NaT: return None
    if isinstance(v, float) and not math.isfinite(v): return None
    return v


def write_json(path, obj):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(path.suffix + ".tmp")
    tmp.write_text(json.dumps(json_safe(obj), indent=2, ensure_ascii=False, allow_nan=False), encoding="utf-8")
    tmp.replace(path)


def sha(path): return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def unique(df, keys, name):
    if df.empty: raise ValueError(f"{name}: no observations")
    if df.duplicated(keys).any(): raise ValueError(f"{name}: duplicate keys {keys}")


def pct(series): return series.pct_change(fill_method=None) * 100


class Sources:
    """Validated raw cache with hashes, retrieval metadata and explicit offline use."""
    def __init__(self, root, args):
        self.root, self.args, self.records = root, args, []
        root.mkdir(parents=True, exist_ok=True)
        self.session = requests.Session()
        self.session.headers["User-Agent"] = "CanadianMultifamilyResearch/2.0 (public data research)"

    def fetch(self, url, name, kind="zip", immutable=False):
        path = self.root / name
        sidecar = path.with_suffix(path.suffix + ".metadata.json")
        metadata = json.loads(sidecar.read_text()) if sidecar.exists() else {}
        valid = path.exists() and metadata.get("sha256") == sha(path) and metadata.get("url") == url
        age = (time.time() - path.stat().st_mtime) / 86400 if path.exists() else np.inf
        cached = valid and (self.args.offline or (not self.args.refresh and (immutable or age < self.args.cache_days)))
        if not cached:
            if self.args.offline:
                raise RuntimeError(f"Offline cache missing or invalid: {name}. Run online once or use --refresh.")
            for attempt in range(3):
                try:
                    LOG.info("Fetching %s", name)
                    response = self.session.get(url, timeout=(15, 120))
                    response.raise_for_status()
                    payload = response.content
                    self.validate(payload, kind)
                    tmp = path.with_suffix(path.suffix + ".download")
                    tmp.write_bytes(payload); tmp.replace(path)
                    metadata = dict(url=url, retrieved_at=datetime.now(timezone.utc).isoformat(),
                                    sha256=sha(path), bytes=len(payload), etag=response.headers.get("ETag"),
                                    last_modified=response.headers.get("Last-Modified"))
                    write_json(sidecar, metadata)
                    break
                except (requests.RequestException, ValueError, zipfile.BadZipFile) as exc:
                    if attempt == 2: raise RuntimeError(f"Download/validation failed: {url}") from exc
                    time.sleep(attempt + 1)
        else:
            LOG.info("Using validated cache %s", name)
            self.validate(path.read_bytes(), kind)
        self.records.append(dict(metadata, file=name, cache_used=cached, offline=self.args.offline))
        return path

    @staticmethod
    def validate(payload, kind):
        if kind in {"zip", "xlsx"}:
            with zipfile.ZipFile(io.BytesIO(payload)) as z:
                if z.testzip(): raise ValueError("Corrupt ZIP member")
                if kind == "xlsx" and "xl/workbook.xml" not in z.namelist(): raise ValueError("Not an XLSX workbook")
                if kind == "zip" and not any(n.endswith(".csv") for n in z.namelist()): raise ValueError("No CSV in ZIP")
        elif kind == "json": json.loads(payload)

    def statcan(self, table, filters):
        path = self.fetch(f"https://www150.statcan.gc.ca/n1/tbl/csv/{table}-eng.zip", f"statcan_{table}.zip")
        kept = []
        with zipfile.ZipFile(path) as z:
            with z.open(f"{table}.csv") as f:
                for chunk in pd.read_csv(f, chunksize=100000, low_memory=False):
                    mask = pd.Series(True, index=chunk.index)
                    for col, values in filters.items(): mask &= chunk[col].isin(values)
                    if mask.any(): kept.append(chunk.loc[mask].copy())
        if not kept: raise ValueError(f"No requested observations in {table}")
        d = pd.concat(kept, ignore_index=True)
        if not d.SCALAR_FACTOR.eq("units").all(): raise ValueError(f"Unexpected scaling in {table}")
        d["source_table"] = table
        return d


def cmhc_number(value):
    s = text(value)
    statuses = {"++": "not_statistically_significant", "**": "suppressed", "-": "not_available",
                "--": "not_available", "..": "not_available", "...": "not_available", "": "not_available"}
    if s in statuses: return np.nan, statuses[s]
    try: return float(s.replace(",", "").replace("$", "").replace("%", "")), "observed"
    except ValueError: raise ValueError(f"Unrecognized CMHC numeric cell: {s!r}")


def header_year(value):
    if isinstance(value, (datetime, pd.Timestamp)): return value.year
    m = re.fullmatch(r"Oct-(\d{2})", text(value), flags=re.I)
    if not m: raise ValueError(f"Unexpected rental year header {value!r}")
    return 2000 + int(m[1])


def centre_header(raw):
    rows = raw.index[raw[0].map(text).eq("Centre")].tolist()
    if len(rows) != 1: raise ValueError("Expected exactly one Centre header")
    return rows[0]


def parse_rental(path, edition):
    sheets = pd.read_excel(path, sheet_name=None, header=None, dtype=object)
    raw = sheets["Table 1.0"]; h = centre_header(raw)
    if raw.shape[1] != 19: raise ValueError("Unrecognized Table 1.0 layout")
    labels = " ".join(text(v) for v in raw.iloc[:h, :].to_numpy().ravel()).lower()
    for label in ["vacancy", "turnover", "fixed sample", "two bedroom"]:
        if label not in labels: raise ValueError(f"Table 1.0 missing {label}")
    if (header_year(raw.iat[h, 1]), header_year(raw.iat[h, 3])) != (edition - 1, edition):
        raise ValueError("Rental edition/header mismatch")
    records = []
    for i in range(h + 1, len(raw)):
        label = text(raw.iat[i, 0]); geo = RMS_LABELS.get(label)
        if geo is None: continue
        for y, cols in [(edition - 1, [1, 6, 11, 15]), (edition, [3, 8, 13, 17])]:
            rec = dict(geography=geo, geography_raw=label, year=y, survey_date=f"{y}-10-01",
                       rental_edition=edition, rental_source_file=path.name,
                       rental_source_sheet="Table 1.0", rental_source_row=i + 1,
                       rental_coverage="Centres 10,000+" if geo == "Canada" else "CMA / CMA part",
                       unit_scope="Privately initiated rental apartments, structures with 3+ units")
            for metric, col in zip(["vacancy_rate", "turnover_rate", "average_2br_rent", "fixed_sample_rent_growth_pct"], cols):
                val, status = cmhc_number(raw.iat[i, col])
                rec[metric], rec[metric + "_status"] = val, status
                rec[metric + "_quality"] = text(raw.iat[i, col + 1])
            rec["vacancy_change_significance"] = text(raw.iat[i, 5]) if y == edition else "not_reported_in_this_edition"
            records.append(rec)
    d = pd.DataFrame(records); unique(d, ["geography", "year"], "Rental edition")
    if set(d.geography) != set(GEOS): raise ValueError(f"Missing rental geographies in edition {edition}")
    # Universe is separately dated and sourced; never substitute rental condos.
    u = sheets["Table 4.1"]; uh = centre_header(u)
    if "universe" not in text(u.iat[uh - 1, 11]).lower() or "RMS" not in text(u.iat[uh, 13]):
        raise ValueError("Unrecognized RMS universe columns")
    universe = []
    for i in range(uh + 1, len(u)):
        geo = RMS_LABELS.get(text(u.iat[i, 0]))
        if geo is None: continue
        val, status = cmhc_number(u.iat[i, 13])
        universe.append(dict(geography=geo, year=edition, rental_universe=val,
                             universe_status=status, universe_source_file=path.name,
                             universe_source_cell=f"Table 4.1!N{i+1}"))
    turnover = []
    if "Table 6.0" in sheets:
        t = sheets["Table 6.0"]; th = centre_header(t)
        if t.shape[1] == 44:
            if "2 Bedroom" not in text(t.iat[th - 2, 11]): raise ValueError("Two-bedroom block moved")
            configs = [(edition - 1, 10, 32), (edition, 12, 34)]
        elif t.shape[1] == 15 and "two bedrooms only" in " ".join(t[0].map(text)).lower():
            configs = [(edition - 1, 2, 9), (edition, 4, 11)]
        else: raise ValueError("Unrecognized turnover table; refusing to guess")
        for y, tc, nc in configs:
            if header_year(t.iat[th, tc]) != y or header_year(t.iat[th, nc]) != y:
                raise ValueError("Turnover dates do not match edition")
            for i in range(th + 1, len(t)):
                geo = RMS_LABELS.get(text(t.iat[i, 0]))
                if geo is None: continue
                rec = dict(geography=geo, year=y, turnover_edition=edition,
                           turnover_source_file=path.name, turnover_source_row=i+1)
                for metric, col in [("turnover_2br_rent", tc), ("non_turnover_2br_rent", nc)]:
                    rec[metric], rec[metric + "_status"] = cmhc_number(t.iat[i, col])
                    rec[metric + "_quality"] = text(t.iat[i, col+1])
                turnover.append(rec)
    return d, pd.DataFrame(universe), pd.DataFrame(turnover)


def rental_history(sources, args):
    values, universes, turnovers = [], [], []
    for y in range(args.start_year, args.rental_year + 1):
        name = f"rmr-canada-{y}-en.xlsx"
        p = sources.fetch(RMS_BASE + name, name, "xlsx", immutable=True)
        d, u, t = parse_rental(p, y)
        values.append(d); universes.append(u)
        if not t.empty: turnovers.append(t)
    d = pd.concat(values).sort_values("rental_edition").drop_duplicates(["geography", "year"], keep="last")
    u = pd.concat(universes); unique(u, ["geography", "year"], "Universe")
    d = d.merge(u, on=["geography", "year"], how="left", validate="one_to_one")
    if turnovers:
        t = pd.concat(turnovers).sort_values("turnover_edition").drop_duplicates(["geography", "year"], keep="last")
        d = d.merge(t, on=["geography", "year"], how="left", validate="one_to_one")
    d = d.sort_values(["geography", "year"]).reset_index(drop=True)
    consecutive = d.groupby("geography").year.diff().eq(1)
    for src, dest in [("average_2br_rent", "average_rent_change_pct"), ("rental_universe", "rental_universe_growth_pct"),
                      ("turnover_2br_rent", "turnover_rent_change_pct")]:
        d[dest] = d.groupby("geography")[src].transform(pct).where(consecutive)
    d["vacancy_change_pp"] = d.groupby("geography").vacancy_rate.diff().where(consecutive)
    d["turnover_premium_pct"] = (d.turnover_2br_rent / d.non_turnover_2br_rent - 1) * 100
    return d[d.year.ge(args.start_year)].reset_index(drop=True)


def population(sources, args):
    d = sources.statcan("17100148", {"GEO": list(POP_LABELS), "Gender": ["Total - gender"], "Age group": ["All ages"]})
    if not d.UOM.eq("Persons").all(): raise ValueError("Population is not in persons")
    d = d.assign(geography=d.GEO.map(POP_LABELS), year=pd.to_numeric(d.REF_DATE), population=pd.to_numeric(d.VALUE))
    d = d[d.year.between(args.start_year - 1, args.rental_year)].sort_values(["geography", "year"])
    unique(d, ["geography", "year"], "CMA population")
    d["population_yoy_pct"] = d.groupby("geography").population.transform(pct).where(d.groupby("geography").year.diff().eq(1))
    d["population_date"] = pd.to_datetime(d.year.astype(str) + "-07-01")
    d["population_boundary"] = "2021 census boundaries (backcast)"
    d["population_status"] = d.STATUS.fillna("")
    d["population_source_geo"] = d.GEO
    d = d[["geography", "year", "population", "population_yoy_pct", "population_date", "population_boundary", "population_status", "population_source_geo", "DGUID", "VECTOR"]]
    q = sources.statcan("17100009", {"GEO": ["Canada"]})
    q = q.assign(date=pd.to_datetime(q.REF_DATE), population=pd.to_numeric(q.VALUE)).sort_values("date")
    q = q[q.date.between(f"{args.start_year-1}-01-01", args.as_of)].copy()
    unique(q, ["date"], "Quarterly national population")
    q["population_yoy_pct"] = q.population.pct_change(4, fill_method=None) * 100
    return d, q[["date", "population", "population_yoy_pct", "STATUS", "VECTOR"]]


def construction(sources, args):
    d = sources.statcan("34100154", {"GEO": list(CON_LABELS), "Type of unit": ["Total units"]})
    d["date"] = pd.to_datetime(d.REF_DATE)
    d = d[d.date.between(f"{args.start_year-1}-01-01", args.as_of)].copy()
    d["geography"] = d.GEO.map(CON_LABELS)
    d["measure"] = d["Housing estimates"].map(CON_MEASURES)
    if d.measure.isna().any(): raise ValueError("Unknown construction measure")
    unique(d, ["geography", "date", "measure"], "Monthly construction")
    monthly = d.pivot(index=["geography", "date"], columns="measure", values="VALUE").reset_index()
    monthly["year"] = monthly.date.dt.year
    monthly["source_table"] = "34-10-0154-01"
    monthly["coverage"] = "CMHC CMA aggregate" 
    monthly.loc[monthly.geography.ne("Canada CMAs"), "coverage"] = "CMA / Ontario part of Ottawa-Gatineau"
    annual = []
    for (geo, year), group in monthly.groupby(["geography", "year"]):
        if set(group.date.dt.month) != set(range(1, 13)): continue
        rec = dict(geography=geo, year=year, construction_months=12, construction_period="calendar_year",
                   construction_coverage=group.coverage.iloc[0], construction_source="34-10-0154-01")
        for c in ["starts", "completions"]:
            if group[c].notna().sum() != 12: raise ValueError(f"Incomplete annual {c}: {geo} {year}")
            rec[c] = group[c].sum(min_count=12)
        rec["units_under_construction"] = group.loc[group.date.dt.month.eq(12), "units_under_construction"].iloc[0]
        rec["construction_stock_date"] = f"{year}-12-31"
        annual.append(rec)
    annual = pd.DataFrame(annual).sort_values(["geography", "year"])
    unique(annual, ["geography", "year"], "Annual construction")
    # Latest common month across all eight source geographies; exact same-month YTD comparisons.
    last = monthly.groupby("geography").date.max().min()
    current_year, month = last.year, last.month
    ytd = []
    for geo in CON_LABELS.values():
        rec = dict(geography=geo, year=current_year, through_month=month,
                   period_end=(last + pd.offsets.MonthEnd(0)).date().isoformat(),
                   period_type="year_to_date" if month < 12 else "annual")
        for y, suffix in [(current_year, ""), (current_year-1, "_prior")]:
            g = monthly[(monthly.geography == geo) & (monthly.year == y) & monthly.date.dt.month.le(month)]
            if set(g.date.dt.month) != set(range(1, month + 1)): raise ValueError(f"Missing YTD months for {geo} {y}")
            for c in ["starts", "completions"]:
                if g[c].notna().sum() != month: raise ValueError(f"Missing YTD {c}: {geo} {y}")
                rec[c + suffix] = g[c].sum(min_count=month)
        for c in ["starts", "completions"]:
            rec[c + "_yoy_pct"] = (rec[c]/rec[c+"_prior"] - 1)*100 if rec[c+"_prior"] > 0 else np.nan
        rec["units_under_construction"] = monthly.loc[(monthly.geography == geo)&(monthly.date == last), "units_under_construction"].iloc[0]
        ytd.append(rec)
    return monthly, annual, pd.DataFrame(ytd), d[["geography", "date", "measure", "VALUE", "STATUS", "VECTOR", "GEO", "DGUID"]]


def rates(sources, args):
    frames = []
    for name, series in BOC_SERIES.items():
        url = f"https://www.bankofcanada.ca/valet/observations/{series}/json?start_date={args.start_year-1}-01-01&end_date={args.as_of}"
        p = sources.fetch(url, f"boc_{name}_{args.start_year-1}_{args.as_of}.json", "json")
        payload = json.loads(p.read_text())
        if series not in payload.get("seriesDetail", {}): raise ValueError(f"BoC metadata missing {series}")
        rows = [{"date": o["d"], name: o.get(series, {}).get("v")} for o in payload.get("observations", [])]
        d = pd.DataFrame(rows)
        d["date"] = pd.to_datetime(d.date); d[name] = pd.to_numeric(d[name], errors="coerce")
        d = d.dropna(subset=[name]); unique(d, ["date"], name)
        frames.append(d.set_index("date"))
    daily = pd.concat(frames, axis=1, sort=True).sort_index().reset_index()
    monthly = []
    for period, group in daily.groupby(daily.date.dt.to_period("M")):
        rec = {"month": str(period), "month_end": period.end_time.normalize(),
               "complete_month": period.end_time.date() < pd.Timestamp(args.as_of).date()}
        for name in BOC_SERIES:
            valid = group.dropna(subset=[name])
            rec[name] = valid[name].iloc[-1] if len(valid) else np.nan
            rec[name+"_observation_date"] = valid.date.iloc[-1] if len(valid) else pd.NaT
        monthly.append(rec)
    monthly = pd.DataFrame(monthly)
    annual = []
    for year, group in daily.groupby(daily.date.dt.year):
        if year >= pd.Timestamp(args.as_of).year: continue
        if group.date.dt.month.nunique() != 12: raise ValueError(f"Incomplete rate year {year}")
        rec = {"year": year}
        for name in BOC_SERIES:
            rec[name+"_daily_mean"] = group[name].mean()
            rec[name+"_observations"] = int(group[name].notna().sum())
        annual.append(rec)
    return daily, monthly, pd.DataFrame(annual)


def build_annual_research_dataset(rental, pop, con, rate_annual):
    d = rental.merge(pop, on=["geography", "year"], how="left", validate="one_to_one")
    d = d.merge(con, on=["geography", "year"], how="left", validate="one_to_one")
    d = d.merge(rate_annual, on="year", how="left", validate="many_to_one")
    d["construction_join_status"] = np.where(d.geography.eq("Canada"),
        "not_applicable: national rental 10,000+ vs construction CMA coverage", "matched")
    d["completions_per_1000_people"] = 1000*d.completions/d.population
    d["rental_universe_growth_status"] = np.where(d.rental_universe_growth_pct.notna(),
        "change_in_reported_stock; boundary/vintage_sensitive", "not_available")
    return d.sort_values(["geography", "year"]).reset_index(drop=True)


def valuation(base_cap):
    rows = []
    for growth in [-5, -2.5, 0, 2.5, 5, 7.5, 10]:
        for bps in range(-100, 101, 25):
            newcap = base_cap + bps/10000
            if newcap <= 0: raise ValueError("Scenario cap rate must be positive")
            rows.append(dict(noi_growth_pct=growth, cap_rate_change_bps=bps,
                             base_cap_rate_pct=base_cap*100, new_cap_rate_pct=newcap*100,
                             implied_value_change_pct=((1+growth/100)*base_cap/newcap-1)*100))
    required = [dict(cap_rate_change_bps=b, base_cap_rate_pct=base_cap*100,
                     new_cap_rate_pct=(base_cap+b/10000)*100,
                     required_noi_growth_pct=(b/10000)/base_cap*100) for b in range(-100, 101, 25)]
    return pd.DataFrame(rows), pd.DataFrame(required)


def validate_data(annual, rental, pop, monthly_con, con, ytd, daily, monthly, val, args):
    checks = []
    def check(name, condition, detail, severity="error"):
        checks.append(dict(check=name, passed=bool(condition), severity=severity, detail=str(detail)))
    expected = pd.MultiIndex.from_product([GEOS, range(args.start_year, args.rental_year+1)])
    actual = pd.MultiIndex.from_frame(annual[["geography", "year"]])
    check("Rental geography-year coverage", expected.difference(actual).empty, f"{len(actual)} rows; {len(expected)} expected")
    check("Unique annual keys", not annual.duplicated(["geography", "year"]).any(), "geography + year")
    check("Population joins complete", annual.population.notna().all() and annual.population_yoy_pct.notna().all(), f"{annual.population.notna().sum()}/{len(annual)}")
    city = annual[annual.geography.ne("Canada")]
    check("CMA construction joins complete", city[["starts", "completions", "units_under_construction"]].notna().all().all(), f"{len(city)} CMA/aggregate rows")
    check("Annual flows contain 12 months", city.construction_months.eq(12).all(), "YTD flows excluded")
    check("National scopes kept distinct", annual.loc[annual.geography.eq("Canada"), "completions"].isna().all(), "10,000+ rental benchmark is not relabeled CMA supply")
    check("Vacancy range", annual.vacancy_rate.between(0, 30).all(), "[0,30]%")
    check("Rent levels positive", annual.average_2br_rent.gt(0).all(), "All averages > 0 CAD")
    check("Population positive", annual.population.gt(0).all(), "All levels > 0 persons")
    check("Construction nonnegative", monthly_con[["starts", "completions", "units_under_construction"]].ge(0).all().all(), "Monthly source values")
    check("Fixed-sample missingness explained", annual.loc[annual.fixed_sample_rent_growth_pct.isna(), "fixed_sample_rent_growth_pct_status"].isin(["not_statistically_significant", "suppressed", "not_available"]).all(), "Source flags retained; no implicit zero")
    check("Latest market indicators available", annual[annual.year.eq(args.rental_year)].turnover_2br_rent.notna().all(), "Latest turnover table")
    for name, frame in [("annual", annual), ("rates", daily), ("valuation", val)]:
        check(f"No infinite values: {name}", not np.isinf(frame.select_dtypes("number").to_numpy()).any(), "NaN is separately explained")
    check("Unique monthly rates", monthly.month.is_unique, "One row per calendar month; per-series dates retained")
    # Independently reproduce every annual construction sum from source monthly rows.
    totals = monthly_con.groupby(["geography", "year"])[["starts", "completions"]].sum(min_count=12)
    joined = con.set_index(["geography", "year"])[["starts", "completions"]]
    check("Construction annual reconciliation", np.allclose(totals.loc[joined.index], joined), "All complete-year flows equal monthly sums")
    recalc = ((1+val.noi_growth_pct/100)/(val.new_cap_rate_pct/100))/(1/(val.base_cap_rate_pct/100))-1
    check("Valuation identity", np.allclose(recalc*100, val.implied_value_change_pct, atol=1e-10), "V = NOI / cap rate")
    check("CMA history sufficient for chart", len(con[con.geography.eq("Canada CMAs") & con.year.between(args.start_year,args.rental_year)]) >= 5, "At least five complete annual points")
    check("Rates reasonably current", (pd.Timestamp(args.as_of)-daily.date.max()).days <= 10, f"Latest source {daily.date.max().date()}")
    check("Population vintage current for rental year", int(pop.year.max()) >= args.rental_year, f"Latest year {pop.year.max()}")
    return pd.DataFrame(checks)


def save_csv(root, name, frame):
    path = root/name; path.parent.mkdir(parents=True, exist_ok=True)
    frame.to_csv(path, index=False, float_format="%.10g", date_format="%Y-%m-%d")


def setup_style():
    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 10, "axes.titlesize": 13,
                         "axes.labelsize": 10, "axes.spines.top": False, "axes.spines.right": False,
                         "axes.edgecolor": "#B4BCC4", "xtick.color": GREY, "ytick.color": GREY,
                         "axes.labelcolor": BLUE, "text.color": BLUE, "figure.facecolor": "white",
                         "savefig.facecolor": "white", "grid.color": "#E4E8EB", "grid.linewidth": .6})


def finish(fig, axes, path, footer, left=.10, right=.97, grid=True):
    for ax in np.array(axes, dtype=object).ravel():
        ax.grid(grid, axis="y"); ax.set_axisbelow(True)
    fig.subplots_adjust(bottom=.23, top=.88, left=left, right=right, hspace=.5, wspace=.30)
    fig.text(.10, .065, footer, fontsize=8, color=GREY, va="bottom")
    fig.savefig(path.with_suffix(".png"), dpi=240)
    fig.savefig(path.with_suffix(".pdf"))
    plt.close(fig)


def charts(root, annual, con, ytd, daily, monthly, sensitivity, args):
    root.mkdir(parents=True, exist_ok=True); setup_style()
    a = annual[annual.geography.eq("Canada")]
    fig, axs = plt.subplots(1, 2, figsize=(10,4.6))
    for ax, c, title in [(axs[0], "vacancy_rate", "Vacancy rate (%)"), (axs[1], "fixed_sample_rent_growth_pct", "Same-sample 2BR rent growth (%)")]:
        ax.plot(a.year, a[c], "o-", color=BLUE, linewidth=2)
        ax.set_title(title, loc="left"); ax.set_xticks(a.year); ax.set_ylim(bottom=0)
        ax.annotate(f"{a[c].iloc[-1]:.1f}%", (a.year.iloc[-1],a[c].iloc[-1]), xytext=(-6,10), textcoords="offset points", ha="right", fontweight="bold")
    finish(fig, axs, root/"01_canadian_rental_conditions", f"CMHC Rental Market Survey, October {args.start_year}–{args.rental_year}; centres 10,000+.\nPrivately initiated rental apartments, 3+ units. Latest available report vintage per observation.")
    c = con[con.geography.eq("Canada CMAs") & con.year.between(args.start_year,args.rental_year)]
    fig, axs = plt.subplots(1,2,figsize=(10,4.6))
    axs[0].plot(c.year,c.starts/1000,"o-",color=BLUE,label="Starts")
    axs[0].plot(c.year,c.completions/1000,"s-",color=TEAL,label="Completions")
    axs[0].set_title("CMA housing flows (000s)",loc="left");axs[0].legend(frameon=False)
    axs[1].plot(c.year,c.units_under_construction/1000,"o-",color=BLUE)
    axs[1].set_title("December construction stock (000s)",loc="left")
    for ax in axs: ax.set_xticks(c.year);ax.set_ylim(bottom=0)
    finish(fig,axs,root/"02_canadian_construction_cycle","Source: CMHC via Statistics Canada 34-10-0154-01; published CMA aggregate, all tenures/types.\nFlows sum 12 months; stock is December. Geography changes with census definitions; not rental-only supply.")
    fig, ax = plt.subplots(figsize=(10,4.6))
    d=daily[daily.date.dt.year.ge(args.start_year)]
    ax.step(d.date,d.overnight_rate,where="post",color=BLUE,lw=1.7,label="Policy target (daily)")
    m=monthly[monthly.complete_month & monthly.month_end.dt.year.ge(args.start_year)]
    for key,col,label in [("goc_5y",TEAL,"5-year GoC"),("goc_10y",ORANGE,"10-year GoC")]:
        ax.plot(m[key+"_observation_date"],m[key],color=col,lw=1.6,label=label+" (month-end)")
    ax.set_ylabel("Percent");ax.set_title("Policy easing and long-term yields",loc="left");ax.legend(frameon=False,loc="upper left",fontsize=8)
    finish(fig,[ax],root/"03_monetary_policy_and_long_term_yields",f"Source: Bank of Canada Valet. Daily policy target; final valid bond observation in each completed month.\nAs of {args.as_of}. Latest partial-month observations are in the rates table, not connected as completed months.")
    pivot=sensitivity.pivot(index="noi_growth_pct",columns="cap_rate_change_bps",values="implied_value_change_pct")
    fig,ax=plt.subplots(figsize=(10,5.6));limit=max(abs(pivot.min().min()),abs(pivot.max().max()))
    im=ax.imshow(pivot,aspect="auto",cmap="RdBu",norm=TwoSlopeNorm(vmin=-limit,vcenter=0,vmax=limit))
    ax.set_xticks(range(len(pivot.columns)),pivot.columns);ax.set_yticks(range(len(pivot.index)),[f"{x:g}" for x in pivot.index])
    for i in range(len(pivot)):
        for j in range(len(pivot.columns)):
            v=pivot.iloc[i,j];ax.text(j,i,f"{v:.1f}",ha="center",va="center",fontsize=9,color="white" if abs(v)>limit*.65 else BLUE)
    ax.set_xlabel("Cap-rate change (basis points)");ax.set_ylabel("NOI growth over scenario horizon (%)")
    ax.set_title(f"Capital-value sensitivity | starting cap rate {args.base_cap_rate*100:.2f}%",loc="left")
    fig.colorbar(im,ax=ax,label="Capital-value change (%)",fraction=.035,pad=.03)
    finish(fig,[ax],root/"04_valuation_sensitivity","Calculation: V = NOI / cap rate. Illustrative unlevered capital-value change; not total return or forecast.\nAssumes a one-year NOI change and end-of-year cap-rate repricing. Excludes interim income, debt, capex and transaction costs.", right=.90, grid=False)
    market=annual[(annual.year==args.rental_year)&annual.geography.isin(CITIES)].set_index("geography").loc[CITIES]
    fig,axs=plt.subplots(1,2,figsize=(10,5))
    labels=[DISPLAY.get(g,g) for g in CITIES]; ys=np.arange(len(CITIES))
    for ax,col,title in [(axs[0],"vacancy_change_pp","Vacancy change (percentage points)"),(axs[1],"turnover_rent_change_pct","Turnover 2BR average-rent change (%)")]:
        vals=market[col].to_numpy();ax.barh(ys,vals,color=[TEAL if v<0 else BLUE for v in vals],height=.58)
        ax.set_yticks(ys,labels if ax is axs[0] else []);ax.invert_yaxis();ax.axvline(0,color=GREY,lw=.6)
        ax.set_title(title,loc="left",fontsize=11);ax.margins(x=.24)
        for y,v in zip(ys,vals):ax.text(v,y,f" {v:+.1f}" if v>=0 else f"{v:+.1f} ",va="center",ha="left" if v>=0 else "right",fontsize=9)
    finish(fig,axs,root/"05_market_conditions",f"Source: CMHC, {args.rental_year-1}–{args.rental_year} October surveys. Cities shown in a fixed order, not ranked.\nTurnover rent compares averages across years, not matched units; vacancy refers to all apartment bedroom types.", left=.22)
    yy=ytd[ytd.geography.eq("Canada CMAs")].iloc[0]
    fig,ax=plt.subplots(figsize=(8,4.6));x=np.arange(2)
    ax.bar(x-.18,[yy.starts_prior/1000,yy.completions_prior/1000],.36,color="#A4B7C5",label=str(int(yy.year)-1))
    ax.bar(x+.18,[yy.starts/1000,yy.completions/1000],.36,color=BLUE,label=str(int(yy.year)))
    ax.set_xticks(x,["Starts","Completions"]);ax.set_ylabel("Units (000s)");ax.legend(frameon=False)
    ax.set_title(f"CMA construction: January–{pd.Timestamp(yy.period_end):%B}",loc="left")
    for bars in ax.containers:ax.bar_label(bars,fmt="%.1f",padding=3,fontsize=9)
    ax.margins(y=.16)
    finish(fig,[ax],root/"06_construction_ytd",f"Source: CMHC via Statistics Canada 34-10-0154-01; same months in each year, through {yy.period_end}.\nPublished CMA aggregate; all dwelling types and tenures. Partial-year flows are not annualized.")
    return market.reset_index()


def display_frame(df):
    return df.replace({"geography": DISPLAY, "fixed_sample_rent_growth_pct_status": {"not_statistically_significant": "Not significant (++)"}}).rename(columns={"geography":"Market", "year":"Year", "vacancy_rate":"Vacancy (%)",
        "vacancy_change_pp":"Vacancy change (pp)", "average_2br_rent":"Average 2BR rent (CAD/month)",
        "fixed_sample_rent_growth_pct":"Same-sample rent growth (%)", "fixed_sample_rent_growth_pct_status":"Rent-growth status",
        "turnover_rent_change_pct":"Turnover average change (%)", "population_yoy_pct":"July population growth (%)",
        "rental_universe_growth_pct":"Rental stock change (%)", "starts":"Housing starts (units)",
        "completions":"Housing completions (units)", "units_under_construction":"Under construction (units)",
        "population":"July population (persons)", "rental_universe":"RMS apartments (units)",
        "completions_per_1000_people":"All-tenure completions / 1,000 people"})


def table_definitions(annual, con, ytd, daily, sensitivity, required, sources, args):
    a=annual[annual.geography.eq("Canada")]
    m=annual[(annual.year==args.rental_year)&annual.geography.isin(CITIES)].set_index("geography").loc[CITIES].reset_index()
    latest=a.iloc[-1]
    snapshot=pd.DataFrame([
        ["Vacancy rate",latest.vacancy_rate,"%",f"October {args.rental_year}","Rental centres 10,000+"],
        ["Vacancy change",latest.vacancy_change_pp,"percentage points",f"Oct {args.rental_year-1} to Oct {args.rental_year}","Rental centres 10,000+"],
        ["Average 2BR rent",latest.average_2br_rent,"CAD/month",f"October {args.rental_year}","New and existing rental apartments"],
        ["Same-sample 2BR rent growth",latest.fixed_sample_rent_growth_pct,"%",f"Oct {args.rental_year-1} to Oct {args.rental_year}","Existing structures, fixed sample"],
        ["July population growth",latest.population_yoy_pct,"%",f"July {args.rental_year-1} to July {args.rental_year}","All Canada; not just rental centres"],
        ["Starting cap-rate assumption",args.base_cap_rate*100,"%","One-year sensitivity","Illustrative, not observed market cap rate"],
    ],columns=["Metric","Value","Unit","Reference period","Scope"])
    national=display_frame(a[["year","vacancy_rate","fixed_sample_rent_growth_pct","average_2br_rent","population_yoy_pct"]])
    markets=display_frame(m[["geography","vacancy_rate","vacancy_change_pp","fixed_sample_rent_growth_pct","fixed_sample_rent_growth_pct_status","turnover_rent_change_pct","population_yoy_pct"]])
    market_supply=display_frame(m[["geography","rental_universe","rental_universe_growth_pct","starts","completions","completions_per_1000_people"]])
    history=con[con.geography.eq("Canada CMAs")&con.year.between(args.start_year,args.rental_year)]
    supply=display_frame(history[["year","starts","completions","units_under_construction"]])
    yt=ytd[["geography","period_end","starts_prior","starts","starts_yoy_pct","completions_prior","completions","completions_yoy_pct"]].copy()
    yt["geography"] = yt["geography"].replace(DISPLAY)
    yt.columns=["Market","YTD through","Prior starts","Current starts","Starts change (%)","Prior completions","Current completions","Completions change (%)"]
    rate_rows=[]
    for name in BOC_SERIES:
        row=daily.dropna(subset=[name]).iloc[-1]
        rate_rows.append([name,row[name],row.date.date().isoformat(),BOC_SERIES[name]])
    rt=pd.DataFrame(rate_rows,columns=["Series","Rate (%)","Observation date","BoC series ID"])
    pivot=sensitivity.pivot(index="noi_growth_pct",columns="cap_rate_change_bps",values="implied_value_change_pct").reset_index()
    pivot.columns=["NOI growth (%)"]+[f"{b:+d} bp" for b in range(-100,101,25)]
    req=required.rename(columns={"cap_rate_change_bps":"Cap-rate change (bp)","base_cap_rate_pct":"Starting cap rate (%)","new_cap_rate_pct":"New cap rate (%)","required_noi_growth_pct":"Break-even NOI growth (%)"})
    methods=pd.DataFrame([
        ["Purpose","Descriptive evidence for a two-page note; no forecast, causal estimate or investment ranking."],
        ["Rental scope","Private rental apartment structures with 3+ units. National headline: centres 10,000+. CMAs separately identified."],
        ["Rent measures","Same-sample growth is primary. Changes in published average/turnover rents are separate composition-sensitive statistics."],
        ["Source flags","++ means statistically indistinguishable from zero, not an exact zero. ** is suppressed. Neither is imputed."],
        ["Geography","Ottawa is Ontario part. CMA population uses 2021 boundaries backcast. CMHC boundaries change with census definitions."],
        ["Periods","October rental surveys; July 1 population; calendar-year construction sums; December stock. Descriptive alignment only."],
        ["Supply","All tenures/types; not rental-only. Ratios use population, not rental stock. CMA aggregate does not equal national 10,000+ coverage."],
        ["Rental stock growth","Changes across annual workbook vintages; boundary/revision-sensitive. Do not interpret as gross completions."],
        ["Rates","Daily-observation annual means, not average month-end points or calendar-time weighting. Per-series dates retained."],
        ["Regression","Not estimated: seven national survey years do not support the proposed multivariable time-series model. No score or ranking."],
        ["Valuation",f"V=NOI/cap rate; starting cap {args.base_cap_rate*100:.2f}%; one-year NOI and end-year repricing; excludes cash yield, leverage, capex, taxes and transaction costs."],
        ["Revisions","Latest available edition per rental observation; exact input URLs, dates and hashes in source_manifest.json. No historical release-date backtest."],
        ["Missing data","Blank numeric fields remain unavailable; source statuses are in clean CSVs. No forward filling of research inputs."],
        ["Reproduce","Run with --offline using the validated raw cache; --refresh redownloads inputs. Reports record code and source hashes."],
    ],columns=["Topic","Definition / limitation"])
    source_table=pd.DataFrame([{"Input":r["file"],"URL":r["url"],"Retrieved (UTC)":r["retrieved_at"],"SHA256":r["sha256"]} for r in sources.records])
    return {
        "National snapshot": (snapshot, "National rental and demand context", "CMHC RMS and Statistics Canada. Periods/scopes are explicitly stated below."),
        "National history": (national, "Rental conditions and population", "October rental observations; July population. CMHC RMS; Statistics Canada 17-10-0148."),
        "Market conditions": (markets, f"Selected rental markets | {args.rental_year}", "Fixed city order, no composite ranking. ++ is not a numerical zero. CMHC RMS; Statistics Canada 17-10-0148."),
        "Market supply": (market_supply, f"Rental stock and all-tenure construction | {args.rental_year}", "October rental stock; calendar-year construction; July population. CMHC RMS; Statistics Canada 34-10-0154 / 17-10-0148."),
        "CMA construction": (supply, "Construction across the CMA aggregate", "All dwelling types/tenures. Annual flows; December stock. CMHC via Statistics Canada 34-10-0154."),
        "Construction YTD": (yt, "Latest construction versus the same months last year", "Same-month cumulative flows, not annualized. CMHC via Statistics Canada 34-10-0154."),
        "Latest rates": (rt, "Latest available policy and bond yields", "Bank of Canada Valet. Each observation is dated; yields are not property cap rates."),
        "Value sensitivity": (pivot, f"Capital-value change (%) | initial cap rate {args.base_cap_rate*100:.2f}%", "Columns: cap-rate change. Rows: one-year NOI growth. Illustrative unlevered capital value, excludes interim income."),
        "Required NOI": (req, "NOI growth required to offset cap-rate repricing", "One-year scenario: required growth = new cap rate / initial cap rate − 1. Assumptions, not forecasts."),
        "Methods": (methods, "Definitions and interpretation", "Read alongside the tables; source quality flags and provenance remain in the clean datasets."),
        "Sources": (source_table, "Input provenance", "Exact downloads used in this run. Hashes permit byte-level verification; see also source_manifest.json."),
    }


def write_workbook(path, tables):
    """Portable formatting in the existing Python/openpyxl output engine."""
    wb=Workbook();wb.remove(wb.active)
    for name,(frame,title,note) in tables.items():
        ws=wb.create_sheet(name); ws.sheet_view.showGridLines=False
        ws.cell(2,2,title).font=Font(name="Arial",size=15,bold=True,color="173F5F")
        ws.cell(3,2,note).font=Font(name="Arial",size=9,italic=True,color="68717A")
        ws.row_dimensions[2].height=25;ws.row_dimensions[3].height=20
        headers=list(frame.columns)
        for j,h in enumerate(headers,2):
            cell=ws.cell(5,j,str(h));cell.fill=PatternFill("solid",fgColor="173F5F")
            cell.font=Font(name="Arial",bold=True,color="FFFFFF",size=10)
            cell.alignment=Alignment(wrap_text=True,vertical="center",horizontal="center")
        ws.row_dimensions[5].height=42
        for i,row in enumerate(frame.itertuples(index=False,name=None),6):
            ws.row_dimensions[i].height=24
            for j,value in enumerate(row,2):
                value=json_safe(value)
                if value is not None and isinstance(value,float) and not math.isfinite(value):value=None
                cell=ws.cell(i,j,value);cell.font=Font(name="Arial",size=10,color="173F5F")
                cell.alignment=Alignment(vertical="center",horizontal="right" if isinstance(value,(int,float)) else "left")
                if i%2==0:cell.fill=PatternFill("solid",fgColor="F1F5F7")
                header=str(headers[j-2]).lower()
                if isinstance(value,(int,float)):
                    if header=="year":cell.number_format="0"
                    elif "population growth" in header or "cap rate" in header or header=="rate (%)":cell.number_format='0.00"%";[Red](0.00"%");0.00"%"'
                    elif "/ 1,000" in header:cell.number_format="0.0"
                    elif "units" in header or "persons" in header or "starts" in header and "change" not in header or "completions" in header and "change" not in header:
                        cell.number_format='#,##0'
                    elif "cad" in header:cell.number_format='"$"#,##0'
                    elif "(%)" in header:cell.number_format='0.0"%";[Red](0.0"%");0.0"%"'
                    elif "(pp)" in header:cell.number_format='0.0" pp"'
                    elif " bp" in header and name=="Value sensitivity":cell.number_format='0.0"%";[Red](0.0"%");0.0"%"'
                    elif "(bp)" in header:cell.number_format='0'
                    else:cell.number_format='0.0'
        ws.column_dimensions['A'].width=3
        for j,h in enumerate(headers,2):
            col=get_column_letter(j)
            if name=="Methods":width=25 if j==2 else 110
            elif name=="Sources":width={2:46,3:85,4:29,5:68}[j]
            elif name=="National snapshot":width={2:36,3:16,4:23,5:29,6:49}[j]
            elif name=="Value sensitivity":width=17 if j==2 else 12
            elif h in ["Market","Series"]:width=25
            elif "status" in str(h).lower():width=34
            else:width=max(17,min(28,len(str(h))*.65))
            ws.column_dimensions[col].width=width
        if name in ["Methods","Sources"]:
            for row in ws.iter_rows(min_row=6):
                for cell in row:cell.alignment=Alignment(wrap_text=True,vertical="top")
                ws.row_dimensions[row[0].row].height=46 if name=="Methods" else 52
        ws.freeze_panes="C6" if len(frame)>12 or len(headers)>6 else "B6"
        ws.auto_filter.ref=f"B5:{get_column_letter(len(headers)+1)}{len(frame)+5}"
        ws.sheet_properties.pageSetUpPr.fitToPage=True
        ws.page_setup.orientation="landscape";ws.page_setup.paperSize=ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
        ws.print_title_rows="1:5"
        ws.print_options.horizontalCentered=True
        ws.print_area=f"B2:{get_column_letter(len(headers)+1)}{len(frame)+5}"
        ws.oddFooter.center.text="Page &P of &N"
        if name=="National snapshot":
            for row in range(6,len(frame)+6):
                unit=ws.cell(row,4).value
                ws.cell(row,3).number_format='"$"#,##0' if unit=="CAD/month" else '0.0'
    wb.save(path)
    # Validate the persisted representation, not only the in-memory workbook.
    check=load_workbook(path,data_only=True)
    if check.sheetnames!=list(tables):raise ValueError("Workbook sheets changed on export")
    for name,(frame,_,_) in tables.items():
        ws=check[name]
        for i,row in enumerate(frame.itertuples(index=False,name=None),6):
            for j,v in enumerate(row,2):
                expected=json_safe(v);actual=ws.cell(i,j).value
                if expected is None:continue
                if isinstance(expected,(int,float)):
                    if not isinstance(actual,(int,float)) or not np.isclose(actual,expected,rtol=1e-12):raise ValueError(f"Workbook numeric mismatch {name}!{i},{j}")
                elif str(actual)!=str(expected):raise ValueError(f"Workbook text mismatch {name}!{i},{j}")
    check.close()


def summaries(root, annual, con, ytd, quarterly, daily, args, checks):
    latest=annual[(annual.year==args.rental_year)&annual.geography.eq("Canada")].iloc[0]
    current=annual[(annual.year==args.rental_year)&annual.geography.isin(CITIES)]
    ys=con[con.geography.eq("Canada CMAs")&con.year.eq(args.rental_year)].iloc[0]
    yp=con[con.geography.eq("Canada CMAs")&con.year.eq(args.rental_year-1)].iloc[0]
    q=quarterly.iloc[-1]
    model={"status":"not_estimated", "national_years":int(annual.year.nunique()),
           "reason":"Short annual history and changing geographic boundaries do not support the proposed multivariable time-series inference.",
           "market_ranking":"not_produced; underlying indicators are shown separately"}
    result={"version":VERSION,"generated_at":datetime.now(timezone.utc).isoformat(),"as_of":args.as_of,
        "rental_year":args.rental_year,"national_rental_scope":"Centres 10,000+; private rental apartments 3+ units",
        "national_rental":latest.to_dict(),"markets":current.to_dict('records'),"construction_ytd":ytd.to_dict('records'),
        "latest_population":q.to_dict(),"model":model,"quality_checks":checks.to_dict('records')}
    write_json(root/"research_summary.json",result)
    declines=current[current.turnover_rent_change_pct.lt(0)].set_index("geography")
    falling=", ".join(f"{DISPLAY.get(g,g)} ({row.turnover_rent_change_pct:+.1f}%)" for g,row in declines.iterrows())
    memo=f'''# Evidence brief for a two-page Canadian multifamily note

Data retrieved as of {args.as_of}; rental survey: October {args.rental_year}. This is a descriptive evidence brief, not a forecast or a statement about any specific investment portfolio.

## Suggested thesis

Rental-market normalization creates a gap between growth in rents paid across the existing stock and pricing on newly turned-over units. Underwriting should distinguish those measures and test NOI against cap-rate repricing.

## Verified facts

- National purpose-built apartment vacancy in centres 10,000+ is {latest.vacancy_rate:.1f}%, up {latest.vacancy_change_pp:.1f} percentage points from the previous October.
- Average two-bedroom rent is ${latest.average_2br_rent:,.0f} per month. Same-sample rent growth is {latest.fixed_sample_rent_growth_pct:.1f}%; the change in the published average is {latest.average_rent_change_pct:.1f}%. The measures differ because the average includes composition changes.
- Average rents for turnover two-bedroom units fell in {falling}. These are changes in group averages, not matched-unit rent indexes.
- In the separately defined CMA construction aggregate, {args.rental_year} completions were {ys.completions:,.0f} ({(ys.completions/yp.completions-1)*100:+.1f}% versus {args.rental_year-1}); starts were {ys.starts:,.0f}. These figures cover all housing tenures and dwelling types.
- The latest national quarterly population estimate is {q.population:,.0f} at {q.date:%B %d, %Y}, with year-over-year growth of {q.population_yoy_pct:.2f}%. City comparisons use July 1 CMA estimates, not province proxies.
- At an illustrative starting cap rate of {args.base_cap_rate*100:.2f}%, +50 bp requires {(0.005/args.base_cap_rate)*100:.1f}% NOI growth to preserve capital value. This excludes interim cash yield, financing, capex and transaction costs.

## Interpretation to develop

Higher vacancy and softer turnover pricing can constrain new-lease income even while same-sample rents continue to rise. Rent growth does not pass one-for-one into NOI: concessions, occupancy and operating costs matter. Long government yields provide financing context, but their movement is not a direct estimate of multifamily cap-rate changes. Supply must be assessed by market and tenure; the all-tenure construction series is contextual evidence.

## Suggested two-page structure

Page 1: the operating-market thesis, national rental history (Chart 01), and selected-market vacancy/turnover comparison (Chart 05). Page 2: supply and financing context, a small valuation sensitivity panel (Chart 04), and implications/risks. Use the formatted workbook for exact numbers; keep the note's data periods visible.

## Essential limitations

October rental, July population and calendar-year construction are different reference periods. Population history uses fixed 2021 boundaries; CMHC construction boundaries change with censuses. The national 10,000+ rental series and construction CMA aggregate are different universes. Calgary's ++ rent-growth flag means not statistically different from zero, not an exact zero. No causal model, estimated investment-return ranking or portfolio-specific conclusion is supported by this package.

## Sources

- CMHC Rental Market Survey archived Canada workbooks ({RMS_PAGE}), Tables 1.0, 4.1 and 6.0.
- Statistics Canada 17-10-0148-01 (July CMA population), 17-10-0009-01 (national quarterly population), and 34-10-0154-01 (CMHC monthly construction).
- Bank of Canada Valet: V39079, BD.CDN.5YR.DQ.YLD, BD.CDN.10YR.DQ.YLD.

Exact input URLs, retrieval timestamps and SHA256 hashes are in diagnostics/source_manifest.json. Tables report source concepts without imputing missing observations.
'''
    (root/"Evidence_Brief.md").write_text(memo,encoding="utf-8")
    write_json(root/"model_status.json",model)


def arguments(argv=None):
    parser=argparse.ArgumentParser(description=__doc__,formatter_class=argparse.RawDescriptionHelpFormatter)
    today=datetime.now(ZoneInfo("America/Toronto")).date()
    parser.add_argument("--output-dir",type=Path,default=Path(__file__).resolve().parent/"output")
    parser.add_argument("--as-of",default=today.isoformat(),help="Retrieval cutoff; not a historical data-vintage backtest")
    parser.add_argument("--start-year",type=int,default=2019)
    parser.add_argument("--rental-year",type=int,default=None)
    parser.add_argument("--base-cap-rate",type=float,default=.05,help="Decimal, e.g. 0.05")
    parser.add_argument("--cache-days",type=float,default=7)
    group=parser.add_mutually_exclusive_group();group.add_argument("--offline",action="store_true");group.add_argument("--refresh",action="store_true")
    args=parser.parse_args(argv)
    cutoff=pd.Timestamp(args.as_of)
    if cutoff.date()>today:parser.error("--as-of cannot be in the future")
    if args.rental_year is None:args.rental_year=cutoff.year-1
    if not 2019<=args.start_year<=args.rental_year:parser.error("Supported archive begins in 2019; check year range")
    if args.rental_year>=cutoff.year:parser.error("Select a completed prior-year rental edition")
    if args.rental_year-args.start_year<4:parser.error("Use at least five rental years for the history figures")
    if args.base_cap_rate<=.01 or args.base_cap_rate>.20:parser.error("Base cap rate must exceed 1% and be at most 20%")
    if args.cache_days<0:parser.error("Cache age cannot be negative")
    return args


def main(argv=None):
    args=arguments(argv);out=args.output_dir.resolve();out.mkdir(parents=True,exist_ok=True)
    run_id=datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%S%fZ")
    staging=out/(".staging-"+run_id);staging.mkdir()
    for sub in ["clean","tables","charts","diagnostics","summary"]:(staging/sub).mkdir()
    LOG.setLevel(logging.INFO);LOG.handlers.clear()
    formatter=logging.Formatter("%(asctime)s | %(levelname)s | %(message)s")
    handler=logging.FileHandler(staging/"diagnostics/pipeline.log",encoding="utf-8");handler.setFormatter(formatter);LOG.addHandler(handler)
    console=logging.StreamHandler();console.setFormatter(formatter);LOG.addHandler(console)
    sources=Sources(out/"raw",args)
    try:
        LOG.info("START %s | version %s | as-of %s",run_id,VERSION,args.as_of)
        LOG.info("1/7: Rental history and source-quality flags")
        rental=rental_history(sources,args)
        LOG.info("2/7: CMA population and latest national estimates")
        pop,quarterly=population(sources,args)
        LOG.info("3/7: Monthly construction, full-year and same-period YTD totals")
        mc,con,ytd,con_long=construction(sources,args)
        LOG.info("4/7: Policy target and government yields")
        daily,monthly,rate_annual=rates(sources,args)
        annual=build_annual_research_dataset(rental,pop,con,rate_annual)
        sensitivity,required=valuation(args.base_cap_rate)
        LOG.info("5/7: Coverage, reconciliation and numeric validation")
        checks=validate_data(annual,rental,pop,mc,con,ytd,daily,monthly,sensitivity,args)
        save_csv(staging/"diagnostics","data_quality_checks.csv",checks)
        failures=checks[~checks.passed & checks.severity.eq("error")]
        if not failures.empty:raise RuntimeError("Validation failed: "+"; ".join(failures.check))
        clean={"cmhc_rental_market.csv":rental,"population_annual.csv":pop,"population_quarterly.csv":quarterly,
               "construction_monthly.csv":mc,"construction_source_observations.csv":con_long,"cmhc_construction.csv":con,
               "construction_ytd.csv":ytd,"boc_daily.csv":daily,"boc_monthly.csv":monthly,"boc_annual.csv":rate_annual,
               "annual_research_dataset.csv":annual}
        for name,frame in clean.items():save_csv(staging/"clean",name,frame)
        LOG.info("6/7: Charts, tables and evidence brief")
        markets=charts(staging/"charts",annual,con,ytd,daily,monthly,sensitivity,args)
        save_csv(staging/"tables","market_conditions.csv",markets)
        save_csv(staging/"tables","valuation_sensitivity.csv",sensitivity)
        save_csv(staging/"tables","required_noi_growth.csv",required)
        tables=table_definitions(annual,con,ytd,daily,sensitivity,required,sources,args)
        write_workbook(staging/"tables/research_tables.xlsx",tables)
        summaries(staging/"summary",annual,con,ytd,quarterly,daily,args,checks)
        write_json(staging/"diagnostics/source_manifest.json",sources.records)
        dictionary=[]
        for name,frame in clean.items():
            for column in frame:
                dictionary.append(dict(file=name,column=column,dtype=str(frame[column].dtype),nonmissing=int(frame[column].notna().sum()),rows=len(frame)))
        save_csv(staging/"diagnostics","data_coverage.csv",pd.DataFrame(dictionary))
        # Strict JSON round-trip and publication artifact inventory.
        for p in staging.rglob("*.json"):
            json.loads(p.read_text(),parse_constant=lambda x: (_ for _ in ()).throw(ValueError(f"Non-standard JSON {x}")))
        charts_found=list((staging/"charts").glob("*.png"))
        if len(charts_found)!=6 or any(p.stat().st_size<15000 for p in charts_found):raise ValueError("Chart export incomplete")
        LOG.info("7/7: Publishing validated outputs (%s checks passed)",len(checks))
        manifest=dict(run_id=run_id,status="complete",version=VERSION,as_of=args.as_of,
                      generated_at=datetime.now(timezone.utc).isoformat(),code_sha256=sha(__file__),
                      arguments={**vars(args), "output_dir": "<runtime output directory>"},python=platform.python_version(),pandas=pd.__version__,
                      numpy=np.__version__,matplotlib=matplotlib.__version__,
                      analytical_scope="Descriptive evidence; no causal regression or composite ranking",
                      outputs=[dict(path=str(p.relative_to(staging)),sha256=sha(p),bytes=p.stat().st_size)
                               for p in staging.rglob("*") if p.is_file() and p.suffix!=".log"])
        write_json(staging/"diagnostics/run_manifest.json",manifest)
        LOG.info("PIPELINE COMPLETE. Output: %s",out)
        handler.close();LOG.removeHandler(handler)
        backup=out/".archive"/run_id
        moved=[]
        try:
            for sub in ["clean","tables","charts","diagnostics","summary"]:
                if (out/sub).exists():backup.mkdir(parents=True,exist_ok=True);(out/sub).rename(backup/sub)
                moved.append(sub);(staging/sub).rename(out/sub)
        except Exception:
            for sub in reversed(moved):
                if (out/sub).exists():shutil.rmtree(out/sub)
                if (backup/sub).exists():(backup/sub).rename(out/sub)
            raise
        staging.rmdir()
        print(f"\nValidated descriptive research package: {out}\nWorkbook: {out/'tables/research_tables.xlsx'}\nEvidence brief: {out/'summary/Evidence_Brief.md'}")
    except Exception:
        LOG.exception("Run failed; previous published outputs remain intact. Diagnostic staging: %s",staging)
        write_json(staging/"diagnostics/run_failure.json",dict(run_id=run_id,status="failed",sources=sources.records))
        raise
    finally:
        sources.session.close()
        for h in list(LOG.handlers):h.close();LOG.removeHandler(h)


if __name__=="__main__":
    main()
